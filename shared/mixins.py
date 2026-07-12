import io
from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import redirect


class StaffRequiredMixin:
    staff_redirect_url = '/'
    staff_error_message = 'You do not have permission to perform this action. Staff access required.'

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_staff:
            messages.error(request, self.staff_error_message)
            return redirect(self.staff_redirect_url)
        return super().dispatch(request, *args, **kwargs)


class ExportMixin:
    """
    Mixin genérico para exportar el queryset filtrado de cualquier ListView.

    Uso en una subclase de ListView:

        export_title  = 'Products'
        export_fields = [
            ('Name',      'name'),           # atributo directo
            ('Brand',     'brand__name'),     # relación FK con doble guión bajo
            ('Suppliers', lambda obj: ', '.join(s.name for s in obj.suppliers.all())),
        ]

    Subclases pueden sobreescribir get_export_fields() para columnas dinámicas.

    El mixin intercepta GET ?format=pdf  → devuelve PDF
                               ?format=xlsx → devuelve Excel
    En cualquier otro caso delega en el ListView normal (con paginación, etc.).
    El export siempre incluye TODOS los registros del queryset filtrado,
    sin importar la página activa.
    """

    export_title = 'Export'
    export_fields = []  # [(header_label, accessor), ...]

    # ------------------------------------------------------------------ #
    #  Helpers                                                             #
    # ------------------------------------------------------------------ #

    def get_export_fields(self):
        """Retorna los campos de exportación activos. Sobreescribir para columnas dinámicas."""
        return self.export_fields

    def _get_value(self, obj, accessor):
        """Resuelve un valor desde un objeto usando string 'a__b__c' o callable."""
        if callable(accessor):
            return accessor(obj)
        value = obj
        for part in accessor.split('__'):
            value = getattr(value, part, '')
            if callable(value):
                value = value()
        return '' if value is None else value

    def _build_rows(self):
        """Retorna lista de listas con los datos del queryset completo."""
        return [
            [str(self._get_value(obj, acc)) for _, acc in self.get_export_fields()]
            for obj in self.get_queryset()
        ]

    # ------------------------------------------------------------------ #
    #  Intercepción de GET                                                 #
    # ------------------------------------------------------------------ #

    def get(self, request, *args, **kwargs):
        fmt = request.GET.get('format', '')
        if fmt == 'pdf':
            return self._export_pdf()
        if fmt == 'xlsx':
            return self._export_xlsx()
        return super().get(request, *args, **kwargs)

    # ------------------------------------------------------------------ #
    #  Exportación PDF                                                     #
    # ------------------------------------------------------------------ #

    def _export_pdf(self):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )

        fields = self.get_export_fields()
        headers = [h for h, _ in fields]
        data = [headers] + self._build_rows()

        num_cols = len(headers)

        # Orientación y fuente dinámicas según cantidad de columnas
        if num_cols <= 3:
            pagesize = A4
            usable_w = A4[0] - 60
            font_size = 9
        elif num_cols <= 5:
            pagesize = landscape(A4)
            usable_w = A4[1] - 60
            font_size = 9
        elif num_cols <= 7:
            pagesize = landscape(A4)
            usable_w = A4[1] - 60
            font_size = 8
        else:
            pagesize = landscape(A4)
            usable_w = A4[1] - 60
            font_size = 7

        col_w = usable_w / max(num_cols, 1)
        col_widths = [col_w] * num_cols

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=pagesize,
            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30,
        )

        dark  = colors.HexColor('#343a40')
        light = colors.HexColor('#f8f9fa')
        grey  = colors.HexColor('#dee2e6')

        style_cmds = [
            ('BACKGROUND',    (0, 0), (-1, 0), dark),
            ('TEXTCOLOR',     (0, 0), (-1, 0), colors.white),
            ('FONTNAME',      (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME',      (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE',      (0, 0), (-1, -1), font_size),
            ('ALIGN',         (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING',    (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING',   (0, 0), (-1, -1), 6),
            ('GRID',          (0, 0), (-1, -1), 0.4, grey),
        ]
        for i in range(2, len(data), 2):
            style_cmds.append(('BACKGROUND', (0, i), (-1, i), light))

        table = Table(data, repeatRows=1, colWidths=col_widths)
        table.setStyle(TableStyle(style_cmds))

        styles = getSampleStyleSheet()
        doc.build([
            Paragraph(self.export_title, styles['Title']),
            Spacer(1, 12),
            table,
        ])
        buf.seek(0)

        resp = HttpResponse(buf, content_type='application/pdf')
        resp['Content-Disposition'] = (
            f'attachment; filename="{self.export_title}.pdf"'
        )
        return resp

    # ------------------------------------------------------------------ #
    #  Exportación Excel                                                   #
    # ------------------------------------------------------------------ #

    def _export_xlsx(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        fields = self.get_export_fields()
        headers = [h for h, _ in fields]
        rows = self._build_rows()

        wb = Workbook()
        ws = wb.active
        ws.title = self.export_title[:31]

        hdr_fill = PatternFill('solid', fgColor='343A40')
        alt_fill = PatternFill('solid', fgColor='F8F9FA')
        hdr_font = Font(bold=True, color='FFFFFF')
        thin_side = Side(style='thin', color='DEE2E6')
        border = Border(
            left=thin_side, right=thin_side,
            top=thin_side, bottom=thin_side,
        )

        ws.append(headers)
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        ws.row_dimensions[1].height = 20

        for r_idx, row in enumerate(rows, start=2):
            ws.append(row)
            use_alt = (r_idx % 2 == 0)
            for cell in ws[r_idx]:
                cell.alignment = Alignment(vertical='center')
                cell.border = border
                if use_alt:
                    cell.fill = alt_fill

        for col_idx, col_cells in enumerate(ws.columns, start=1):
            max_len = max(
                (len(str(c.value or '')) for c in col_cells),
                default=8,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(
                max_len + 4, 50
            )

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        resp = HttpResponse(
            buf,
            content_type=(
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
        )
        resp['Content-Disposition'] = (
            f'attachment; filename="{self.export_title}.xlsx"'
        )
        return resp

class GroupRequiredMixin:
    """
    Mixin que verifica si el usuario pertenece a alguno
    de los roles (grupos) indicados en group_required.

    Uso:
        class GroupListView(LoginRequiredMixin, GroupRequiredMixin, ListView):
            group_required = ['Administrador']
    """
    group_required = []        # Lista de roles permitidos
    group_redirect_url = '/'   # A dónde redirigir si no tiene el rol
    group_error_message = 'You do not have permission to access this option.'

    def dispatch(self, request, *args, **kwargs):
        # 1. Si no inició sesión -> al login
        if not request.user.is_authenticated:
            return redirect('security:login')
        # 2. El superusuario siempre pasa
        if request.user.is_superuser:
            return super().dispatch(request, *args, **kwargs)
        # 3. ¿Pertenece a alguno de los roles permitidos?
        if request.user.groups.filter(name__in=self.group_required).exists():
            return super().dispatch(request, *args, **kwargs)
        # 4. No tiene el rol -> mensaje de error y redirección
        messages.error(request, self.group_error_message)
        return redirect(self.group_redirect_url)